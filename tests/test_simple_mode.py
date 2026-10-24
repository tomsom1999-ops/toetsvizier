import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from toetsanalyse.simple_model import (create_analysis, save_as, open_analysis, set_counts,
    active_questions, active_students, completion, project_active, import_cito)
from toetsanalyse.results import save_score, stored_results, test_questions
from toetsanalyse.norming import save_normalization, has_active_normalization
from toetsanalyse.analysis import analysis_data


class SimpleModeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db = create_analysis(3, "Testanalyse", 2)

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def fill(self):
        with self.db.connection:
            self.db.connection.execute("UPDATE matrix_questions SET maximum_score=2 WHERE test_id=1")
            for student in active_students(self.db):
                for q in active_questions(self.db):
                    save_score(self.db, 1, student["id"], q["id"], "N" if q["label"] == "3" else "0" if q["label"] == "2" else "2", commit=False)

    def qt(self):
        try:
            from PySide6.QtWidgets import QApplication
        except ModuleNotFoundError:
            self.skipTest("PySide6 ontbreekt")
        self.app = QApplication.instance() or QApplication([])

    def test_counts_hide_and_restore_data_including_saved_file(self):
        self.fill()
        old = stored_results(self.db, 1)
        with self.db.connection:
            set_counts(self.db, 1, 1)
        self.assertEqual(1, len(active_questions(self.db)))
        self.assertEqual(1, len(active_students(self.db)))
        self.assertEqual(old, stored_results(self.db, 1))
        path = Path(self.temp.name) / "analyse.tviz"
        save_as(self.db, path)
        reopened = open_analysis(path)
        try:
            with reopened.connection:
                set_counts(reopened, 3, 2)
            self.assertEqual(old, stored_results(reopened, 1))
            self.assertEqual(3, len(active_questions(reopened)))
            self.assertEqual(2, len(active_students(reopened)))
        finally:
            reopened.close()

    def test_create_flow_chooses_file_first_and_opens_an_already_saved_analysis(self):
        self.qt()
        from PySide6.QtWidgets import QMainWindow, QDialog, QSpinBox
        from toetsanalyse.app import MainWindow
        from unittest.mock import Mock
        host = QMainWindow()
        host.show_simple = Mock()
        path = Path(self.temp.name) / 'SE2 Trillingen.tviz'
        order = []
        def choose(*args, **kwargs):
            order.append('bestand')
            return str(path), ''
        def counts(dialog):
            order.append('aantallen')
            self.assertEqual(['bestand', 'aantallen'], order)
            counters = dialog.findChildren(QSpinBox)
            counters[0].setValue(3)
            counters[1].setValue(2)
            return QDialog.DialogCode.Accepted
        try:
            with patch('toetsanalyse.app.QFileDialog.getSaveFileName', side_effect=choose), \
                 patch('toetsanalyse.app.QDialog.exec', new=counts):
                MainWindow.create_simple(host)
            host.show_simple.assert_called_once()
            created = host.show_simple.call_args.args[0]
            try:
                self.assertTrue(path.exists())
                self.assertEqual(path, created.path)
                self.assertEqual('SE2 Trillingen', created.scalar('SELECT name FROM tests WHERE id=1'))
                reopened = open_analysis(path)
                try:
                    self.assertEqual(3, len(active_questions(reopened)))
                    self.assertEqual(2, len(active_students(reopened)))
                finally:
                    reopened.close()
            finally:
                created.close()
        finally:
            host.close()

    def test_cancelling_new_file_or_counts_does_not_create_a_document(self):
        self.qt()
        from PySide6.QtWidgets import QMainWindow, QDialog
        from toetsanalyse.app import MainWindow
        from unittest.mock import Mock
        host = QMainWindow()
        host.show_simple = Mock()
        try:
            with patch('toetsanalyse.app.QFileDialog.getSaveFileName', return_value=('', '')), \
                 patch('toetsanalyse.app.QDialog.exec') as counts:
                MainWindow.create_simple(host)
                counts.assert_not_called()
            path = Path(self.temp.name) / 'cancelled.tviz'
            with patch('toetsanalyse.app.QFileDialog.getSaveFileName', return_value=(str(path), '')), \
                 patch('toetsanalyse.app.QDialog.exec', return_value=QDialog.DialogCode.Rejected):
                MainWindow.create_simple(host)
            self.assertFalse(path.exists())
            host.show_simple.assert_not_called()
        finally:
            host.close()

    def test_cito_flow_chooses_destination_before_import_and_saves_before_opening(self):
        self.qt()
        from PySide6.QtWidgets import QMainWindow
        from toetsanalyse.app import MainWindow
        from unittest.mock import Mock
        host = QMainWindow()
        host.show_simple = Mock()
        imported = create_analysis(3, 'Oorspronkelijke CITO-naam', 2)
        path = Path(self.temp.name) / 'Mijn CITO-analyse.tviz'
        try:
            with patch('toetsanalyse.app.QFileDialog.getOpenFileName', return_value=('cito.xlsx', '')), \
                 patch('toetsanalyse.app.QFileDialog.getSaveFileName', return_value=('', '')), \
                 patch('toetsanalyse.simple_model.import_cito') as parser:
                MainWindow.import_cito_simple(host)
                parser.assert_not_called()
                host.show_simple.assert_not_called()
            with patch('toetsanalyse.app.QFileDialog.getOpenFileName', return_value=('cito.xlsx', '')), \
                 patch('toetsanalyse.app.QFileDialog.getSaveFileName', return_value=(str(path), '')), \
                 patch('toetsanalyse.simple_model.import_cito', return_value=imported):
                MainWindow.import_cito_simple(host)
            self.assertTrue(path.exists())
            self.assertEqual('Mijn CITO-analyse', imported.scalar('SELECT name FROM tests WHERE id=1'))
            self.assertIs(imported, host.show_simple.call_args.args[0])
        finally:
            imported.close()
            host.close()

    def test_opening_older_file_replaces_placeholder_title_with_filename(self):
        path = Path(self.temp.name) / 'Bestaande toets.tviz'
        save_as(self.db, path)
        with self.db.connection:
            self.db.connection.execute("UPDATE tests SET name='Nieuwe analyse' WHERE id=1")
            self.db.set_meta('subject_name', 'Nieuwe analyse')
        reopened = open_analysis(path)
        try:
            self.assertEqual('Bestaande toets', reopened.scalar('SELECT name FROM tests WHERE id=1'))
            self.assertEqual('Bestaande toets', reopened.meta('subject_name'))
        finally:
            reopened.close()

    def test_saved_filename_becomes_title_and_edits_are_written_without_save(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QMainWindow
        from toetsanalyse.app import MainWindow
        from toetsanalyse.simple_workspace import SimpleWorkspace
        original = Path(self.temp.name) / 'SE2 Trillingen.tviz'
        save_as(self.db, original)
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'autosave.ini'), QSettings.Format.IniFormat))
        host = QMainWindow()
        w.title_changed.connect(lambda name: MainWindow.set_simple_title(host, name))
        try:
            self.assertEqual('SE2 Trillingen', w.name.text())
            w.edit_cells([(1, 1, '1')])
            reopened = open_analysis(original)
            try:
                sid = active_students(reopened)[0]['id']
                qid = active_questions(reopened)[0]['id']
                self.assertEqual(1, stored_results(reopened, 1)[1][sid, qid])
                self.assertEqual('SE2 Trillingen', reopened.scalar('SELECT name FROM tests WHERE id=1'))
            finally:
                reopened.close()
            target = Path(self.temp.name) / 'Nieuwe kopie.tviz'
            with patch('toetsanalyse.simple_workspace.QFileDialog.getSaveFileName', return_value=(str(target), '')):
                self.assertTrue(w.save(True))
            self.assertEqual('Nieuwe kopie', w.name.text())
            self.assertTrue(host.windowTitle().startswith('Nieuwe kopie · ToetsVizier'))
            self.assertEqual('Nieuwe kopie', w.windowTitle())
            w.name.setText('Mijn aangepaste titel')
            w.name.editingFinished.emit()
            w.save()
            reopened = open_analysis(target)
            try:
                self.assertEqual('Mijn aangepaste titel', reopened.scalar('SELECT name FROM tests WHERE id=1'))
            finally:
                reopened.close()
            reopened = open_analysis(original)
            try:
                self.assertEqual('SE2 Trillingen', reopened.scalar('SELECT name FROM tests WHERE id=1'))
            finally:
                reopened.close()
        finally:
            w.close()
            host.close()

    def test_projection_only_analyses_active_data_and_persists_norm(self):
        self.fill()
        with self.db.connection:
            set_counts(self.db, 1, 1)
        projected = project_active(self.db)
        try:
            result = analysis_data(projected, 1)
            self.assertEqual(1, result["participant_count"])
            self.assertEqual(1, result["question_count"])
            self.assertEqual(2, result["maximum_score"])
            self.assertEqual(2, result["participants"][0]["total_score"])
            save_normalization(projected, 1, {"method": "n_term", "n_term": 1.0})
            self.assertTrue(has_active_normalization(self.db, 1))
        finally:
            projected.close()
        self.assertEqual(3, len(test_questions(self.db, 1)))
        self.assertEqual(6, self.db.scalar("SELECT COUNT(*) FROM scores"))

    def test_modules_require_all_active_scores_but_zero_and_n_count(self):
        self.assertFalse(completion(self.db)["ready"])
        self.fill()
        self.assertTrue(completion(self.db)["ready"])
        sid = active_students(self.db)[0]["id"]
        qid = active_questions(self.db)[0]["id"]
        save_score(self.db, 1, sid, qid, "")
        self.assertFalse(completion(self.db)["ready"])
        self.assertEqual(1, len(completion(self.db)["missing"]))

    def test_wrong_file_kind_is_rejected_before_any_migration(self):
        from toetsanalyse.database import SubjectDatabase
        path = Path(self.temp.name) / "vak.db"
        other = SubjectDatabase.create(path, "Vak", "2026")
        other.close()
        original = path.read_bytes()
        with self.assertRaises(ValueError):
            open_analysis(path)
        self.assertEqual(original, path.read_bytes())

    def test_clipboard_names_scores_zero_n_undo_and_autosave(self):
        self.qt()
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QApplication
        from toetsanalyse.simple_workspace import SimpleWorkspace
        settings = QSettings(str(Path(self.temp.name) / "settings.ini"), QSettings.Format.IniFormat)
        workspace = SimpleWorkspace(self.db, settings)
        workspace.table.setCurrentCell(0, 1)
        QApplication.clipboard().setText("2\t2\t2")
        workspace.paste()
        workspace.table.setCurrentCell(1, 0)
        QApplication.clipboard().setText("Sam\t2\t0\tN\nAlex\t1\t2\t2")
        workspace.paste()
        self.assertTrue(workspace.norm_button.isEnabled())
        self.assertTrue(workspace.analysis_button.isEnabled())
        self.assertEqual("Sam", active_students(self.db)[0]["display_name"])
        path = Path(self.temp.name) / "saved.tviz"
        save_as(self.db, path)
        workspace.edit_cells([(1, 1, "1")])
        workspace.undo()
        reopened = open_analysis(path)
        try:
            sid = active_students(reopened)[0]["id"]
            qid = active_questions(reopened)[0]["id"]
            self.assertEqual(2, stored_results(reopened, 1)[1][sid, qid])
        finally:
            reopened.close()
        workspace.question_count.setValue(1)
        workspace.student_count.setValue(1)
        workspace.change_counts()
        self.assertEqual(1, len(workspace.questions))
        workspace.question_count.setValue(3)
        workspace.student_count.setValue(2)
        workspace.change_counts()
        self.assertEqual("N", workspace.table.item(1, 3).text())
        self.assertEqual("Alex", workspace.table.item(2, 0).text())
        workspace.close()

    def test_invalid_paste_is_atomic_and_blocks_modules_until_corrected(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings
        from toetsanalyse.simple_workspace import SimpleWorkspace
        workspace = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / "s.ini"), QSettings.Format.IniFormat))
        before = stored_results(self.db, 1)
        self.assertFalse(workspace.edit_cells([(1, 1, "1"), (1, 2, "9")]))
        self.assertEqual(before, stored_results(self.db, 1))
        self.assertFalse(workspace.analysis_button.isEnabled())
        self.assertTrue(workspace.undo_button.isEnabled())
        workspace.undo()
        self.assertTrue(workspace.analysis_button.isEnabled())
        workspace.close()

    def test_metadata_categories_and_multiple_choice_survive_hidden_columns(self):
        self.qt()
        from toetsanalyse.simple_dialogs import MetadataDialog
        self.fill()
        dialog = MetadataDialog(self.db)
        dialog.definitions.append({"name": "Mijn indeling", "standard": False, "values": ["Een", "Twee"]})
        dialog.enabled.add('Mijn indeling')
        dialog.build()
        qid = active_questions(self.db)[2]["id"]
        dialog.combos[qid, "Mijn indeling"].setCurrentText("Twee")
        with self.db.connection:
            dialog.apply()
            self.db.connection.execute("UPDATE matrix_questions SET is_multiple_choice=1,multiple_choice_answer='A' WHERE id=?", (qid,))
            set_counts(self.db, 1, 1)
        visible = MetadataDialog(self.db)
        with self.db.connection:
            visible.apply()
            set_counts(self.db, 3, 2)
        self.assertEqual("Twee", self.db.scalar("SELECT v.name FROM question_taxonomy_values q JOIN taxonomy_values v ON v.id=q.taxonomy_value_id WHERE q.question_id=?", (qid,)))
        self.assertEqual("A", self.db.scalar("SELECT multiple_choice_answer FROM matrix_questions WHERE id=?", (qid,)))
        dialog.close()
        visible.close()

    def test_metadata_checkboxes_control_analysis_export_and_preserve_unchecked_values(self):
        self.qt()
        self.fill()
        from openpyxl import load_workbook
        from toetsanalyse.simple_dialogs import MetadataDialog
        from toetsanalyse.simple_model import active_taxonomies
        from toetsanalyse.simple_workspace import export_simple_excel
        qid = active_questions(self.db)[0]['id']
        dialog = MetadataDialog(self.db)
        self.assertTrue({'RTTI', 'OBIT', 'Bloom'}.issubset(dialog.checks))
        self.assertEqual(set(), dialog.enabled)
        dialog.checks['RTTI'].setChecked(True)
        dialog.combos[qid, 'RTTI'].setCurrentIndex(1)
        rtti = dialog.combos[qid, 'RTTI'].currentText()
        dialog.checks['OBIT'].setChecked(True)
        dialog.combos[qid, 'OBIT'].setCurrentIndex(1)
        obit = dialog.combos[qid, 'OBIT'].currentText()
        dialog.checks['RTTI'].setChecked(False)
        self.assertNotIn((qid, 'RTTI'), dialog.combos)
        dialog.checks['RTTI'].setChecked(True)
        self.assertEqual(rtti, dialog.combos[qid, 'RTTI'].currentText())
        dialog.checks['OBIT'].setChecked(False)
        with self.db.connection:
            dialog.apply()
        self.assertEqual(['RTTI'], [row['name'] for row in active_taxonomies(self.db)])
        source_count = self.db.scalar('SELECT COUNT(*) FROM question_taxonomy_values')
        projected = project_active(self.db)
        try:
            titles = {d['title'] for d in analysis_data(projected, 1)['group_dimensions']}
            self.assertIn('RTTI', titles)
            self.assertNotIn('OBIT', titles)
        finally:
            projected.close()
        self.assertEqual(source_count, self.db.scalar('SELECT COUNT(*) FROM question_taxonomy_values'))
        path = Path(self.temp.name) / 'selected.xlsx'
        export_simple_excel(self.db, path)
        workbook = load_workbook(path)
        try:
            header = next(workbook['Vragen en metadata'].values)
            self.assertIn('RTTI', header)
            self.assertNotIn('OBIT', header)
        finally:
            workbook.close()
        save_as(self.db, Path(self.temp.name) / 'selected.tviz')
        reopened = open_analysis(self.db.path)
        try:
            saved_dialog = MetadataDialog(reopened)
            self.assertEqual({'RTTI'}, saved_dialog.enabled)
            saved_dialog.checks['RTTI'].setChecked(False)
            with reopened.connection:
                saved_dialog.apply()
            empty = MetadataDialog(reopened)
            self.assertEqual(set(), empty.enabled)
            before = reopened.connection.serialize()
            empty.checks['OBIT'].setChecked(True)
            self.assertEqual(obit, empty.combos[qid, 'OBIT'].currentText())
            self.assertEqual(before, reopened.connection.serialize())
            empty.reject()
            self.assertEqual(before, reopened.connection.serialize())
            saved_dialog.close()
            empty.close()
        finally:
            reopened.close()
        dialog.close()

    def test_question_header_shows_mc_key_and_buttons_have_distinct_roles(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QPushButton
        from toetsanalyse.simple_workspace import SimpleWorkspace
        qid = active_questions(self.db)[1]['id']
        with self.db.connection:
            self.db.connection.execute("UPDATE matrix_questions SET is_multiple_choice=1,multiple_choice_answer='B' WHERE id=?", (qid,))
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'buttons.ini'), QSettings.Format.IniFormat))
        try:
            self.assertEqual('Vraag 2\nSleutel: B', w.table.horizontalHeaderItem(2).text())
            self.assertEqual('Vraag 1', w.table.horizontalHeaderItem(1).text())
            self.assertGreaterEqual(w.table.horizontalHeader().height(), 54)
            buttons = {b.text(): b for b in w.findChildren(QPushButton)}
            for text in ['Opslaan als', 'Exporteren naar Excel', 'Uitleg', 'Sluiten', 'Vragen indelen',
                         'Meerkeuze instellen', 'Ongedaan maken', 'Groot invoerscherm']:
                self.assertEqual('secondaryButton', buttons[text].objectName())
            self.assertEqual('primaryButton', buttons['Opslaan'].objectName())
            self.assertEqual('dangerButton', w.remove_button.objectName())
            w.table.setCurrentCell(0, 1)
            self.assertFalse(w.remove_button.isEnabled())
            w.table.setCurrentCell(1, 1)
            self.assertTrue(w.remove_button.isEnabled())
            with self.db.connection:
                self.db.connection.execute("UPDATE matrix_questions SET multiple_choice_answer='C' WHERE id=?", (qid,))
            w.refresh()
            self.assertIn('Sleutel: C', w.table.horizontalHeaderItem(2).text())
        finally:
            w.close()

    def test_multiple_custom_distributions_can_be_added_together_and_reopened(self):
        self.qt()
        self.fill()
        from toetsanalyse.simple_dialogs import MetadataDialog
        dialog = MetadataDialog(self.db)
        qid = active_questions(self.db)[0]['id']
        try:
            with patch('toetsanalyse.simple_dialogs.text_input', side_effect=['Hoofdstuk', 'H1\nH2']):
                dialog.add_distribution()
            dialog.combos[qid, 'Hoofdstuk'].setCurrentText('H1')
            with patch('toetsanalyse.simple_dialogs.text_input', side_effect=['Leerdoel', 'Doel A\nDoel B']):
                dialog.add_distribution()
            self.assertEqual('H1', dialog.combos[qid, 'Hoofdstuk'].currentText())
            dialog.combos[qid, 'Leerdoel'].setCurrentText('Doel B')
            self.assertTrue(dialog.checks['Hoofdstuk'].isChecked())
            self.assertTrue(dialog.checks['Leerdoel'].isChecked())
            with self.db.connection:
                dialog.apply()
            save_as(self.db, Path(self.temp.name) / 'two-distributions.tviz')
            reopened = open_analysis(self.db.path)
            try:
                saved = MetadataDialog(reopened)
                self.assertEqual({'Hoofdstuk', 'Leerdoel'}, saved.enabled)
                self.assertEqual('H1', saved.combos[qid, 'Hoofdstuk'].currentText())
                self.assertEqual('Doel B', saved.combos[qid, 'Leerdoel'].currentText())
                projected = project_active(reopened)
                try:
                    titles = {d['title'] for d in analysis_data(projected, 1)['group_dimensions']}
                    self.assertTrue({'Hoofdstuk', 'Leerdoel'}.issubset(titles))
                finally:
                    projected.close()
                saved.close()
            finally:
                reopened.close()
        finally:
            dialog.close()

    def test_delete_custom_distribution_removes_hidden_links_keeps_scores_and_supports_undo(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QMessageBox, QDialog
        from toetsanalyse.simple_dialogs import MetadataDialog
        from toetsanalyse.simple_workspace import SimpleWorkspace
        qs = active_questions(self.db)
        setup = MetadataDialog(self.db)
        for name in ['Verkeerd', 'Bewaren']:
            with patch('toetsanalyse.simple_dialogs.text_input', side_effect=[name, 'Een\nTwee']):
                setup.add_distribution()
            setup.combos[qs[2]['id'], name].setCurrentText('Twee')
        with self.db.connection:
            setup.apply()
            set_counts(self.db, 1, 1)
        setup.close()
        previous_scores = stored_results(self.db, 1)
        tid = self.db.scalar("SELECT id FROM taxonomy_definitions WHERE name='Verkeerd'")
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'remove.ini'), QSettings.Format.IniFormat))
        dialog = MetadataDialog(self.db, w)
        try:
            with patch('toetsanalyse.simple_dialogs.text_input', return_value='Verkeerd'), \
                 patch('toetsanalyse.simple_dialogs.question', return_value=QMessageBox.StandardButton.Yes):
                dialog.remove_distribution()
            self.assertNotIn('Verkeerd', dialog.checks)
            self.assertIsNotNone(self.db.scalar('SELECT id FROM taxonomy_definitions WHERE id=?', (tid,)))
            with patch('toetsanalyse.simple_workspace.MetadataDialog', return_value=dialog), \
                 patch.object(dialog, 'exec', return_value=QDialog.DialogCode.Accepted):
                w.metadata()
            self.assertIsNone(self.db.scalar('SELECT id FROM taxonomy_definitions WHERE id=?', (tid,)))
            self.assertEqual(0, self.db.scalar('SELECT COUNT(*) FROM question_taxonomy_values WHERE taxonomy_id=?', (tid,)))
            self.assertEqual(0, self.db.scalar('SELECT COUNT(*) FROM taxonomy_values WHERE taxonomy_id=?', (tid,)))
            self.assertEqual(0, self.db.scalar('SELECT COUNT(*) FROM test_taxonomy_selections WHERE taxonomy_id=?', (tid,)))
            self.assertEqual(previous_scores, stored_results(self.db, 1))
            self.assertIsNotNone(self.db.scalar("SELECT id FROM taxonomy_definitions WHERE name='Bewaren'"))
            self.assertEqual(3, self.db.scalar('SELECT COUNT(*) FROM taxonomy_definitions WHERE is_standard=1'))
            w.undo()
            self.assertEqual(tid, self.db.scalar("SELECT id FROM taxonomy_definitions WHERE name='Verkeerd'"))
            self.assertEqual(1, self.db.scalar('SELECT COUNT(*) FROM question_taxonomy_values WHERE taxonomy_id=? AND question_id=?', (tid, qs[2]['id'])))
            self.assertEqual(previous_scores, stored_results(self.db, 1))
        finally:
            dialog.close()
            w.close()

    def test_deleting_distribution_can_be_declined_or_cancelled_and_never_deletes_standards(self):
        self.qt()
        from PySide6.QtWidgets import QMessageBox
        from toetsanalyse.simple_dialogs import MetadataDialog
        dialog = MetadataDialog(self.db)
        self.assertFalse(dialog.remove_distribution_button.isEnabled())
        with patch('toetsanalyse.simple_dialogs.text_input', side_effect=['Eigen', 'Een\nTwee']):
            dialog.add_distribution()
        with self.db.connection:
            dialog.apply()
        before = self.db.connection.serialize()
        with patch('toetsanalyse.simple_dialogs.text_input', return_value='Eigen'), \
             patch('toetsanalyse.simple_dialogs.question', return_value=QMessageBox.StandardButton.No):
            dialog.remove_distribution()
        self.assertIn('Eigen', dialog.checks)
        with patch('toetsanalyse.simple_dialogs.text_input', return_value='RTTI'), \
             patch('toetsanalyse.simple_dialogs.question') as confirm:
            dialog.remove_distribution()
            confirm.assert_not_called()
        with patch('toetsanalyse.simple_dialogs.text_input', return_value='Eigen'), \
             patch('toetsanalyse.simple_dialogs.question', return_value=QMessageBox.StandardButton.Yes):
            dialog.remove_distribution()
        dialog.reject()
        self.assertEqual(before, self.db.connection.serialize())
        dialog.close()

    def test_dynamic_cito_import_skips_unused_rows_and_keeps_blanks(self):
        try:
            from openpyxl import Workbook
        except ModuleNotFoundError:
            self.skipTest("openpyxl ontbreekt")
        workbook = Workbook()
        reference = workbook.active
        reference.title = "Referentie"
        reference["C8"] = 2
        reference["C10"] = "hele"
        data = workbook.create_sheet("Gegevens")
        data["E4"] = "Max. score:"
        for col in (6, 7):
            data.cell(2, col, col - 5)
            data.cell(3, col, f'=IF(1=1,{col - 5},"")')
            data.cell(4, col, 3)
        data["B6"], data["F6"], data["G6"] = "Sam", 0, 3
        data["B7"] = "Afwezig"
        data["B8"], data["F8"] = "Alex", 1
        data["B107"] = "Gem. Toetsscore:"
        path = Path(self.temp.name) / "CITO.xlsx"
        workbook.save(path)
        workbook.close()
        original = path.read_bytes()
        imported = import_cito(path)
        try:
            self.assertEqual(2, len(active_questions(imported)))
            self.assertEqual(2, len(active_students(imported)))
            self.assertEqual(1, len(completion(imported)["missing"]))
            self.assertEqual(0, min(stored_results(imported, 1)[1].values()))
        finally:
            imported.close()
        self.assertEqual(original, path.read_bytes())

    def test_excel_export_only_contains_active_data_and_preserves_letters(self):
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError:
            self.skipTest("openpyxl ontbreekt")
        from toetsanalyse.simple_workspace import export_simple_excel
        self.fill()
        with self.db.connection:
            set_counts(self.db, 2, 1)
        path = Path(self.temp.name) / "scores.xlsx"
        export_simple_excel(self.db, path)
        workbook = load_workbook(path, data_only=True)
        try:
            rows = list(workbook["Scores"].values)
            self.assertEqual(3, len(rows))
            self.assertEqual(4, len(rows[0]))
            self.assertEqual([2, 0, 2], list(rows[2][1:]))
        finally:
            workbook.close()

    def test_native_enter_navigation_horizontal_and_vertical(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings, Qt
        from PySide6.QtTest import QTest
        from toetsanalyse.simple_workspace import SimpleWorkspace
        workspace = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / "keys.ini"), QSettings.Format.IniFormat))
        workspace.show()
        workspace.precision.setCurrentIndex(workspace.precision.findData(.5))
        self.app.processEvents()
        for mode, expected in (("horizontal", (1, 2)), ("vertical", (2, 1))):
            workspace.direction.setCurrentIndex(workspace.direction.findData(mode))
            workspace.table.setCurrentCell(1, 1)
            workspace.table.setFocus()
            QTest.keyClick(workspace.table, Qt.Key.Key_F2)
            QTest.keyClicks(workspace.table.focusWidget(), "1")
            QTest.keyClick(workspace.table.focusWidget(), Qt.Key.Key_Return)
            self.app.processEvents()
            self.assertEqual(expected, (workspace.table.currentRow(), workspace.table.currentColumn()))
        workspace.direction.setCurrentIndex(workspace.direction.findData("hybrid"))
        workspace.block_size.setValue(2)
        workspace.advance(1, 3)
        self.assertEqual((2, 3), (workspace.table.currentRow(), workspace.table.currentColumn()))
        workspace.close()

    def test_whole_scores_advance_per_question_up_to_nine_and_confirm_above_nine(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings, Qt
        from PySide6.QtTest import QTest
        from toetsanalyse.simple_workspace import SimpleWorkspace
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'quick.ini'), QSettings.Format.IniFormat))
        w.show()
        self.app.processEvents()
        try:
            for mode, expected in [('horizontal', (1, 2)), ('vertical', (2, 1)), ('hybrid', (1, 2))]:
                w.direction.setCurrentIndex(w.direction.findData(mode))
                w.table.setCurrentCell(1, 1)
                w.table.setFocus()
                QTest.keyClicks(w.table, '1')
                self.app.processEvents()
                self.assertEqual(expected, (w.table.currentRow(), w.table.currentColumn()))
                self.assertEqual('1', w.table.item(1, 1).text())
            w.direction.setCurrentIndex(w.direction.findData('horizontal'))
            w.edit_cells([(0, 1, '12')])
            w.table.setCurrentCell(1, 1)
            QTest.keyClicks(w.table, '1')
            self.app.processEvents()
            self.assertEqual((1, 1), (w.table.currentRow(), w.table.currentColumn()))
            QTest.keyClicks(w.table.focusWidget(), '2')
            self.app.processEvents()
            self.assertEqual((1, 1), (w.table.currentRow(), w.table.currentColumn()))
            QTest.keyClick(w.table.focusWidget(), Qt.Key.Key_Return)
            self.app.processEvents()
            self.assertEqual('12', w.table.item(1, 1).text())
            self.assertEqual((1, 2), (w.table.currentRow(), w.table.currentColumn()))
            QTest.keyClicks(w.table, '2')
            self.app.processEvents()
            self.assertEqual((1, 3), (w.table.currentRow(), w.table.currentColumn()))
            w.edit_cells([(0, 3, '9')])
            w.table.setCurrentCell(1, 3)
            QTest.keyClicks(w.table, '9')
            self.app.processEvents()
            self.assertEqual('9', w.table.item(1, 3).text())
            self.assertNotEqual((1, 3), (w.table.currentRow(), w.table.currentColumn()))
            w.table.setCurrentCell(1, 3)
            QTest.keyClick(w.table, Qt.Key.Key_F2)
            QTest.keyClicks(w.table.focusWidget(), '8')
            self.app.processEvents()
            self.assertEqual('8', w.table.item(1, 3).text())
            self.assertNotEqual((1, 3), (w.table.currentRow(), w.table.currentColumn()))
            w.precision.setCurrentIndex(w.precision.findData(.5))
            w.table.setCurrentCell(1, 2)
            QTest.keyClicks(w.table, '1')
            self.app.processEvents()
            self.assertEqual((1, 2), (w.table.currentRow(), w.table.currentColumn()))
            QTest.keyClick(w.table.focusWidget(), Qt.Key.Key_Escape)
        finally:
            w.close()

    def test_input_colors_match_extended_and_highlight_active_headers(self):
        self.qt()
        from PySide6.QtCore import QSettings
        from toetsanalyse.simple_workspace import SimpleWorkspace
        questions = active_questions(self.db)
        students = active_students(self.db)
        with self.db.connection:
            self.db.connection.execute('UPDATE matrix_questions SET maximum_score=2 WHERE test_id=1')
            self.db.connection.execute("UPDATE matrix_questions SET is_multiple_choice=1, multiple_choice_answer='B' WHERE id=?", (questions[1]['id'],))
            for student, values in zip(students, [('0', 'A', 'N'), ('1', 'B', '2')]):
                for q, raw in zip(questions, values):
                    save_score(self.db, 1, student['id'], q['id'], raw, commit=False)
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'colors.ini'), QSettings.Format.IniFormat))
        try:
            color = lambda r, c: w.table.item(r, c).background().color().name()
            self.assertEqual('#ffffff', color(1, 1))  # A numeric zero is not a wrong answer.
            self.assertEqual('#fdecec', color(1, 2))
            self.assertEqual('#e8f7e8', color(2, 2))
            self.assertEqual('#fff8ed', color(1, 3))
            self.assertEqual('#dbeafe', color(0, 1))
            self.assertEqual('#dbeafe', color(1, 4))
            self.assertTrue(w.table.item(1, 4).font().bold())
            w.table.setCurrentCell(1, 2)
            self.assertEqual(1, w.table.verticalHeader().active_row)
            self.assertEqual(2, w.table.horizontalHeader().active_column)
            self.assertEqual('#dbeafe', color(1, 0))
            w.table.setCurrentCell(2, 3)
            self.assertEqual('#ffffff', color(1, 0))
            self.assertEqual('#dbeafe', color(2, 0))
            w.clear_scores([(1, 1)])
            self.assertEqual('#edf2f7', color(1, 4))
            self.assertFalse(w.table.item(1, 4).font().bold())
            self.assertFalse(w.edit_cells([(1, 1, '3')]))
            self.assertIn((1, 1), w.invalid)
            self.assertTrue(w.table.item(1, 1).toolTip())
            self.assertIn('rode rand', w.legend.text())
        finally:
            w.close()

    def test_vertical_snake_route_native_confirmation_and_final_cell(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings, Qt
        from PySide6.QtTest import QTest
        from toetsanalyse.simple_workspace import SimpleWorkspace
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'snake.ini'), QSettings.Format.IniFormat))
        w.show()
        self.app.processEvents()
        try:
            self.assertFalse(w.vertical_direction.isVisible())
            w.direction.setCurrentIndex(w.direction.findData('vertical'))
            self.assertTrue(w.vertical_direction.isVisible())
            self.assertGreaterEqual(w.vertical_direction.findText('Omhoog'), 0)
            self.assertGreaterEqual(w.vertical_direction.findText('Omlaag'), 0)
            w.table.setCurrentCell(2, 1)
            w.table.setFocus()
            QTest.keyClick(w.table, Qt.Key.Key_F2)
            QTest.keyClick(w.table.focusWidget(), Qt.Key.Key_Return)  # Unchanged score: still advance exactly once.
            self.assertEqual((2, 2), (w.table.currentRow(), w.table.currentColumn()))
            self.assertEqual('up', w.vertical_direction.currentData())
            for key, expected in zip([Qt.Key.Key_Tab, Qt.Key.Key_Return, Qt.Key.Key_Tab, Qt.Key.Key_Return],
                                     [(1, 2), (1, 3), (2, 3), (2, 3)]):
                QTest.keyClick(w.table, key)
                self.assertEqual(expected, (w.table.currentRow(), w.table.currentColumn()))
            w.direction.setCurrentIndex(w.direction.findData('horizontal'))
            w.table.setCurrentCell(1, 3)
            w.advance(1, 3)
            self.assertEqual((2, 1), (w.table.currentRow(), w.table.currentColumn()))
            w.advance(2, 3)
            self.assertEqual((2, 1), (w.table.currentRow(), w.table.currentColumn()))
        finally:
            w.close()

    def test_hybrid_begin_question_snake_blocks_and_hidden_count_change(self):
        self.qt()
        from PySide6.QtCore import QSettings
        from toetsanalyse.simple_workspace import SimpleWorkspace
        with self.db.connection:
            set_counts(self.db, 5, 2)
        self.fill()
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'blocks.ini'), QSettings.Format.IniFormat))
        try:
            w.direction.setCurrentIndex(w.direction.findData('hybrid'))
            w.block_size.setValue(2)
            w.begin_question.setCurrentIndex(w.begin_question.findData(2))
            w.table.setCurrentCell(1, 2)
            for expected in [(1, 3), (2, 2), (2, 3), (2, 4), (2, 5), (1, 4), (1, 5), (1, 5)]:
                w.advance(w.table.currentRow(), w.table.currentColumn())
                self.assertEqual(expected, (w.table.currentRow(), w.table.currentColumn()))
            self.assertEqual(4, w.begin_question.currentData())
            self.assertEqual('up', w.vertical_direction.currentData())
            w.question_count.setValue(2)
            w.change_counts()
            self.assertEqual(1, w.begin_question.currentData())
            self.assertEqual(2, w.begin_question.count())
        finally:
            w.close()

    def test_direct_n_letters_backspace_and_progress(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings, Qt
        from PySide6.QtTest import QTest
        from toetsanalyse.simple_workspace import SimpleWorkspace
        qs = active_questions(self.db)
        with self.db.connection:
            self.db.connection.execute('DELETE FROM scores WHERE question_id=?', (qs[1]['id'],))
            self.db.connection.execute("UPDATE matrix_questions SET is_multiple_choice=1,multiple_choice_answer='B' WHERE id=?", (qs[1]['id'],))
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'letters.ini'), QSettings.Format.IniFormat))
        w.show()
        self.app.processEvents()
        try:
            self.assertIn('4 van 6', w.progress.text())
            w.precision.setCurrentIndex(w.precision.findData(.5))
            w.table.setCurrentCell(1, 1)
            w.table.setFocus()
            QTest.keyClicks(w.table, 'n')
            self.assertEqual('N', w.table.item(1, 1).text())
            self.assertEqual((1, 2), (w.table.currentRow(), w.table.currentColumn()))
            QTest.keyClicks(w.table, 'b')
            self.assertEqual('B', w.table.item(1, 2).text())
            self.assertEqual('#e8f7e8', w.table.item(1, 2).background().color().name())
            self.assertIn('5 van 6', w.progress.text())
            w.table.setCurrentCell(1, 2)
            QTest.keyClick(w.table, Qt.Key.Key_Backspace)
            self.assertEqual('', w.table.item(1, 2).text())
            self.assertIn('4 van 6', w.progress.text())
            w.undo()
            self.assertEqual('B', w.table.item(1, 2).text())
            w.table.setCurrentCell(1, 2)
            QTest.keyClick(w.table, Qt.Key.Key_F2)
            QTest.keyClick(w.table.focusWidget(), Qt.Key.Key_Backspace)
            self.assertEqual('', w.table.item(1, 2).text())
            w.edit_cells([(1, 1, '9')])
            self.assertIn('3 van 6', w.progress.text())
        finally:
            w.close()

    def test_large_input_reuses_table_shortcuts_state_and_returns_without_data_loss(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings, QTimer, Qt
        from PySide6.QtTest import QTest
        from toetsanalyse.simple_workspace import SimpleWorkspace
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'large.ini'), QSettings.Format.IniFormat))
        w.show()
        self.app.processEvents()
        table = w.table
        errors = []
        def inspect_dialog():
            try:
                self.assertIsNotNone(w.large_dialog)
                self.assertTrue(w.input.isVisible())
                self.assertTrue(w.table.isVisible())
                self.assertIs(w.table, table)
                self.assertTrue(w.large_dialog.isAncestorOf(table))
                self.assertFalse(w.large_dialog.isAncestorOf(w.question_count))
                self.assertFalse(w.large_dialog.isAncestorOf(w.metadata_button))
                self.assertFalse(w.large_dialog.isAncestorOf(w.mc_button))
                from PySide6.QtWidgets import QPushButton
                large_button_texts = {button.text() for button in w.large_dialog.findChildren(QPushButton)}
                self.assertNotIn('Opslaan', large_button_texts)
                self.assertIn('Uitleg invoerscherm', large_button_texts)
                self.assertIn('Terug naar hoofdscherm', large_button_texts)
                w.direction.setCurrentIndex(w.direction.findData('vertical'))
                w.vertical_direction.setCurrentIndex(w.vertical_direction.findData('up'))
                w.edit_cells([(1, 1, '1')], advance=False)
                w.large_dialog.activateWindow()
                table.setFocus()
                self.app.processEvents()
                QTest.qWait(20)
                QTest.keyClick(table, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
                self.app.processEvents()
                self.assertEqual('2', table.item(1, 1).text())
                w.table.setCurrentCell(2, 1)
                from PySide6.QtWidgets import QApplication
                QApplication.clipboard().setText('1')
                w.large_dialog.activateWindow()
                table.setFocus()
                self.app.processEvents()
                QTest.keyClick(table, Qt.Key.Key_V, Qt.KeyboardModifier.ControlModifier)
                self.app.processEvents()
                self.assertEqual('1', table.item(2, 1).text())
                w.undo()
                self.assertEqual('2', table.item(2, 1).text())
                w.edit_cells([(1, 1, '1')], advance=False)
                self.assertEqual(w.saved.text(), w.large_saved.text())
            except Exception as error:
                errors.append(error)
            finally:
                w.large_dialog.accept()
        QTimer.singleShot(50, inspect_dialog)
        try:
            w.open_large_input()
            self.assertEqual([], errors)
            self.assertIs(w.stack.currentWidget(), w.input)
            self.assertIsNone(w.large_dialog)
            self.assertEqual('1', table.item(1, 1).text())
            self.assertEqual('vertical', w.direction.currentData())
            self.assertEqual('up', w.vertical_direction.currentData())
            self.assertTrue(all(s.isEnabled() for s in w.shortcuts))
            self.assertTrue(w.large_button.isVisible())
        finally:
            w.close()

    def test_delete_clears_scores_atomically_preserves_names_maxima_and_can_undo(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings, Qt
        from PySide6.QtTest import QTest
        from PySide6.QtWidgets import QTableWidgetSelectionRange
        from toetsanalyse.simple_workspace import SimpleWorkspace
        path = Path(self.temp.name) / 'delete.tviz'
        save_as(self.db, path)
        w = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / 'delete.ini'), QSettings.Format.IniFormat))
        w.show()
        self.app.processEvents()
        try:
            original = stored_results(self.db, 1)
            name = w.table.item(1, 0).text()
            w.table.setCurrentCell(1, 1)
            w.table.setRangeSelected(QTableWidgetSelectionRange(0, 0, 2, 3), True)
            QTest.keyClick(w.table, Qt.Key.Key_Delete)
            self.app.processEvents()
            self.assertEqual(name, w.table.item(1, 0).text())
            self.assertEqual('2', w.table.item(0, 1).text())
            self.assertEqual('', w.table.item(1, 1).text())
            self.assertEqual('', w.table.item(2, 3).text())
            self.assertFalse(w.analysis_button.isEnabled())
            reopened = open_analysis(path)
            try:
                self.assertFalse(completion(reopened)['ready'])
            finally:
                reopened.close()
            w.undo()
            self.assertEqual(original, stored_results(self.db, 1))
            w.table.setCurrentCell(1, 1)
            QTest.keyClick(w.table, Qt.Key.Key_F2)
            QTest.keyClick(w.table.focusWidget(), Qt.Key.Key_Delete)
            self.app.processEvents()
            self.assertEqual('', w.table.item(1, 1).text())
            self.assertEqual((1, 1), (w.table.currentRow(), w.table.currentColumn()))
        finally:
            w.close()

    def test_open_modules_with_hidden_data_and_return_to_input(self):
        self.qt()
        self.fill()
        from PySide6.QtCore import QSettings
        from toetsanalyse.simple_workspace import SimpleWorkspace
        workspace = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / "modules.ini"), QSettings.Format.IniFormat))
        workspace.question_count.setValue(1)
        workspace.student_count.setValue(1)
        workspace.change_counts()
        for module in ("norm", "analysis"):
            workspace.open_module(module)
            self.assertIsNot(workspace.input, workspace.stack.currentWidget())
            self.assertEqual(1, len(test_questions(workspace.projection, 1)))
            self.assertEqual(1, analysis_data(workspace.projection, 1)["participant_count"])
            workspace.show_input()
        workspace.release_views()
        workspace.close()

    def test_analysis_is_orange_until_norming_and_updates_immediately(self):
        self.qt()
        self.fill()
        import json
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QMessageBox
        from toetsanalyse.app import NormingBridge
        from toetsanalyse.simple_workspace import SimpleWorkspace
        from toetsanalyse.simple_presentation import simplify_dashboard
        workspace = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / "norm-state.ini"), QSettings.Format.IniFormat))
        self.assertTrue(workspace.analysis_button.isEnabled())
        self.assertTrue(workspace.analysis_button.property("partial_analysis"))
        self.assertIn("#d97706", workspace.analysis_button.styleSheet())
        self.assertIn("cijferafhankelijke", workspace.status.text())
        workspace.open_module("analysis")
        data = analysis_data(workspace.projection, 1)
        self.assertFalse(data["normalization_finalized"])
        self.assertTrue(all(p["grade"] is None for p in data["participants"]))
        workspace.open_module("norm")
        bridge = NormingBridge(workspace.active_page)
        result = json.loads(bridge.finalizeNormalization(json.dumps({"method": "n_term", "n_term": 1.0})))
        self.assertTrue(result["ok"])
        self.assertFalse(workspace.analysis_button.property("partial_analysis"))
        self.assertEqual("", workspace.analysis_button.styleSheet())
        with patch("toetsanalyse.simple_widgets.question", return_value=QMessageBox.StandardButton.Yes):
            self.assertTrue(json.loads(bridge.removeNormalization())["ok"])
        self.assertTrue(workspace.analysis_button.property("partial_analysis"))
        without_norm = simplify_dashboard("<head></head><body></body>", "analysis")
        with_norm = simplify_dashboard("<head></head><body></body>", "analysis", normalization_finalized=True)
        self.assertIn('<div class="simple-guide simple-no-norm">', without_norm)
        self.assertNotIn('<div class="simple-guide simple-no-norm">', with_norm)
        self.assertIn("Groepsniveau gaat over de deelnemers", without_norm)
        workspace.release_views()
        workspace.close()

    def test_guided_help_topics_examples_navigation_and_skip(self):
        self.qt()
        from PySide6.QtWidgets import QDialog
        from toetsanalyse.simple_help import InputIntroduction, TOPICS
        guide = InputIntroduction()
        guide.show()
        self.app.processEvents()
        self.assertEqual(7, guide.topic.count())
        self.assertIn('Sleutel: B', guide.demo.table.horizontalHeaderItem(3).text())
        self.assertEqual('#e8f7e8', guide.demo.table.item(1, 3).background().color().name())
        self.assertEqual('#fdecec', guide.demo.table.item(2, 3).background().color().name())
        self.assertFalse(guide.previous.isEnabled())
        guide.next_button.click()
        self.assertEqual(1, guide.topic.currentIndex())
        guide.topic.setCurrentIndex(3)
        self.app.processEvents()
        scores_start = guide.demo.highlights()[0].x()
        guide.variant.setCurrentIndex(1)
        self.assertLess(guide.demo.highlights()[0].x(), scores_start)
        guide.topic.setCurrentIndex(4)
        guide.variant.setCurrentIndex(1)
        self.assertEqual('Verticaal (per vraag)', guide.demo.route_choice.currentText())
        start, end = guide.demo.arrows()[0]
        self.assertEqual(start.x(), end.x())
        self.assertGreater(end.y(), start.y())
        guide.topic.setCurrentIndex(6)
        self.assertIn("automatisch", guide.text.text())
        self.assertTrue(guide.next_button.isEnabled())
        self.assertFalse(guide.skip.isChecked())
        guide.skip_button.click()
        self.assertEqual(QDialog.DialogCode.Accepted, guide.result())
        self.assertFalse(guide.skip.isChecked())

    def test_hybrid_toolbar_fits_small_window(self):
        self.qt()
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QLabel, QPushButton
        from toetsanalyse.app import STYLESHEET
        from toetsanalyse.simple_workspace import SimpleWorkspace
        previous_style = self.app.styleSheet()
        self.app.setStyleSheet(STYLESHEET)
        workspace = SimpleWorkspace(self.db, QSettings(str(Path(self.temp.name) / "layout.ini"), QSettings.Format.IniFormat))
        try:
            workspace.direction.setCurrentIndex(workspace.direction.findData("hybrid"))
            workspace.resize(980, 640)
            workspace.show()
            self.app.processEvents()
            self.assertLessEqual(workspace.width(), 980)
            self.assertLessEqual(workspace.minimumSizeHint().width(), 980)
            labels = {label.text() for label in workspace.input.findChildren(QLabel)}
            self.assertIn('Toets instellen', labels)
            self.assertIn('Invoerroute:', labels)
            self.assertNotIn('Toets- en invoerinstellingen', labels)
            self.assertNotIn('Lager aantal? Extra gegevens blijven bewaard.', labels)
            for control in (workspace.direction, workspace.vertical_direction, workspace.begin_question,
                            workspace.block_size, workspace.precision):
                self.assertGreater(workspace.table.geometry().top(), control.geometry().bottom())
            self.assertGreater(workspace.progress.geometry().top(), workspace.table.geometry().bottom())
            self.assertGreaterEqual(workspace.table.height(), 220)
            for button in workspace.input.findChildren(QPushButton):
                if button.isVisible():
                    self.assertLessEqual(button.mapTo(workspace, button.rect().topRight()).x(), workspace.width())
        finally:
            workspace.close()
            self.app.setStyleSheet(previous_style)

    def test_settings_table_controls_fit_rows_and_use_dutch_buttons(self):
        self.qt()
        from PySide6.QtWidgets import QDialogButtonBox, QTableWidget
        from toetsanalyse.app import STYLESHEET
        from toetsanalyse.simple_dialogs import MetadataDialog, MultipleChoiceDialog
        previous_style = self.app.styleSheet()
        self.app.setStyleSheet(STYLESHEET)
        try:
            for dialog in (MetadataDialog(self.db), MultipleChoiceDialog(self.db)):
                dialog.show()
                self.app.processEvents()
                table = dialog.findChild(QTableWidget)
                for row in range(table.rowCount()):
                    for column in range(table.columnCount()):
                        control = table.cellWidget(row, column)
                        if control is not None:
                            self.assertGreaterEqual(table.rowHeight(row), control.minimumSizeHint().height() + 1)
                buttons = dialog.findChild(QDialogButtonBox)
                self.assertEqual("Opslaan", buttons.button(QDialogButtonBox.StandardButton.Save).text())
                self.assertEqual("Annuleren", buttons.button(QDialogButtonBox.StandardButton.Cancel).text())
                dialog.close()
        finally:
            self.app.setStyleSheet(previous_style)

    def test_simple_options_and_explanations_only_appear_after_opening_mode(self):
        self.qt()
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QPushButton, QLabel
        from PySide6.QtTest import QSignalSpy
        from toetsanalyse.app import StartPage
        start = StartPage(QSettings(str(Path(self.temp.name) / "start.ini"), QSettings.Format.IniFormat))
        buttons = {button.text(): button for button in start.findChildren(QPushButton)}
        for name in ("Simpel openen", "CITO-analyseformulier uploaden", "Zelf starten", "Analyse openen", "Nieuw vak aanmaken", "Bestaand vak openen"):
            self.assertIn(name, buttons)
        start.show()
        self.app.processEvents()
        labels = {label.text(): label for label in start.findChildren(QLabel)}
        self.assertLess(labels['<h2>Simpel</h2>'].mapTo(start, labels['<h2>Simpel</h2>'].rect().topLeft()).y(),
                        labels['<h2>Uitgebreid</h2>'].mapTo(start, labels['<h2>Uitgebreid</h2>'].rect().topLeft()).y())
        self.assertTrue(buttons["Simpel openen"].isVisible())
        for name in ("CITO-analyseformulier uploaden", "Zelf starten", "Analyse openen"):
            self.assertFalse(buttons[name].isVisible())
        buttons["Simpel openen"].click()
        self.app.processEvents()
        self.assertFalse(buttons["Bestaand vak openen"].isVisible())
        for name in ("CITO-analyseformulier uploaden", "Zelf starten", "Analyse openen"):
            self.assertTrue(buttons[name].isVisible())
        spy = QSignalSpy(start.new_simple)
        buttons["Zelf starten"].click()
        self.assertEqual(1, spy.count())
        self.assertTrue(any("analyseformulier met de leerlingnamen" in label.text() for label in start.simple_menu.findChildren(QLabel)))
        buttons["Terug naar het startscherm"].click()
        self.assertTrue(buttons["Simpel openen"].isVisible())
        start.close()


if __name__ == "__main__":
    unittest.main()
