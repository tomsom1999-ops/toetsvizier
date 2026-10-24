from __future__ import annotations

import csv
import io
import sqlite3
from pathlib import Path

from PySide6.QtCore import QEvent, Qt, Signal, QTimer
from PySide6.QtGui import QColor, QKeySequence, QShortcut, QPen
from PySide6.QtWidgets import (QApplication, QCheckBox, QComboBox, QDialog, QFileDialog,
    QFrame, QHBoxLayout, QLabel, QLineEdit, QMessageBox, QPushButton, QSpinBox,
    QStackedWidget, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget, QStyledItemDelegate,
    QAbstractItemDelegate)

from .results import save_score, stored_results, test_questions
from .simple_model import (SIMPLE_SUFFIX, add_student, completion, invalidate_grades,
                           recalculate_grades, save_as, set_counts, active_questions, active_students, project_active)
from .simple_dialogs import InputIntroduction, MetadataDialog, MultipleChoiceDialog
from .simple_widgets import WrappingLayout, question


class ScoreEntryDelegate(QStyledItemDelegate):
    def __init__(self, workspace):
        super().__init__(workspace.table)
        self.workspace = workspace

    def paint(self, painter, option, index):
        super().paint(painter, option, index)
        invalid = (index.row(), index.column()) in self.workspace.invalid
        active = index == self.workspace.table.currentIndex()
        if invalid or active:
            painter.save()
            painter.setPen(QPen(QColor('#dc2626' if invalid else '#2563eb'), 2))
            painter.drawRect(option.rect.adjusted(1, 1, -2, -2))
            painter.restore()

    def createEditor(self, parent, option, index):
        editor = super().createEditor(parent, option, index)
        if isinstance(editor, QLineEdit):
            row, col = index.row(), index.column()
            def edited(text):
                if self.workspace.quick_score_complete(row, col, text):
                    QTimer.singleShot(0, lambda: self.finish_quick_score(editor, row, col))
            editor.textEdited.connect(edited)
        return editor

    def finish_quick_score(self, editor, row, col):
        from shiboken6 import isValid
        if (not isValid(editor) or not editor.isVisible()
                or not self.workspace.quick_score_complete(row, col, editor.text())):
            return
        self.commitData.emit(editor)
        if isValid(editor):
            self.closeEditor.emit(editor, QAbstractItemDelegate.EndEditHint.NoHint)

    def eventFilter(self, editor, event):
        if event.type() == QEvent.Type.KeyPress and event.key() == Qt.Key.Key_Delete:
            row, col = self.workspace.table.currentRow(), self.workspace.table.currentColumn()
            if self.workspace.is_score_cell(row, col):
                self.workspace.clear_scores([(row, col)])
                return True
        if event.type() == QEvent.Type.KeyPress and event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter, Qt.Key.Key_Tab):
            row, col = self.workspace.table.currentRow(), self.workspace.table.currentColumn()
            self.commitData.emit(editor)
            try:
                self.closeEditor.emit(editor, QAbstractItemDelegate.EndEditHint.NoHint)
            except RuntimeError:
                pass  # Refreshing the table can already have closed the editor.
            if (row, col) not in self.workspace.invalid:
                self.workspace.advance(row, col)
            return True
        return super().eventFilter(editor, event)


class SimpleWorkspace(QWidget):
    back_to_start = Signal()

    def __init__(self, database, settings, parent=None):
        super().__init__(parent)
        self.database, self.settings = database, settings
        self.history = []
        self.invalid = {}
        self.loading = False
        self._released = False
        self.active_page = None
        self.projection = None
        root = QVBoxLayout(self)
        top = WrappingLayout()
        top.addWidget(QLabel("Simpel · één losse toets"))
        self.name = QLineEdit(str(database.scalar("SELECT name FROM tests WHERE id=1")))
        self.name.setPlaceholderText("Naam van de analyse")
        self.name.editingFinished.connect(self.rename)
        self.name.setMinimumWidth(150)
        top.addWidget(self.name)
        for label, method in (("Opslaan", self.save), ("Opslaan als", lambda: self.save(True)),
                              ("Exporteren naar Excel", self.export_excel), ("Uitleg", lambda: self.show_intro(True)),
                              ("Sluiten", self.request_close)):
            button = QPushButton(label)
            button.clicked.connect(lambda _=False, fn=method: fn())
            top.addWidget(button)
        root.addLayout(top)
        steps = QHBoxLayout()
        self.input_button = QPushButton("1. Gegevens invullen")
        self.input_button.clicked.connect(self.show_input)
        self.norm_button = QPushButton("2. Normeren")
        self.norm_button.clicked.connect(lambda: self.open_module("norm"))
        self.analysis_button = QPushButton("3. Analyseren")
        self.analysis_button.clicked.connect(lambda: self.open_module("analysis"))
        for button in (self.input_button, self.norm_button, self.analysis_button):
            steps.addWidget(button)
        root.addLayout(steps)
        self.status = QLabel()
        self.status.setWordWrap(True)
        root.addWidget(self.status)
        self.stack = QStackedWidget()
        root.addWidget(self.stack, 1)
        self.input = QWidget()
        layout = QVBoxLayout(self.input)
        hint = QLabel("Vul eerst maximumscores in. Typ leerlingnamen en scores, of kopieer cellen uit Excel en plak met Ctrl+V. "
                      "0 = geen punten · N = niet gemaakt · leeg = nog invullen. Del wist geselecteerde scores. "
                      "Bij hele punten en maximaal 9 punten voor de vraag springt u automatisch verder. "
                      "Afwezige leerlingen neemt u niet op.")
        hint.setWordWrap(True)
        layout.addWidget(hint)
        actions = WrappingLayout()
        self.question_count = QSpinBox()
        self.question_count.setRange(1, 1000)
        self.question_count.setPrefix("Aantal vragen: ")
        self.question_count.setValue(int(database.meta("simple_question_count", "1")))
        self.student_count = QSpinBox()
        self.student_count.setRange(0, 10000)
        self.student_count.setPrefix("Aantal leerlingen: ")
        self.student_count.setValue(int(database.meta("simple_student_count", "0")))
        for counter in (self.question_count, self.student_count):
            counter.setToolTip("Een lager aantal verbergt gegevens. Verhoog het aantal weer om ze terug te zien.")
            counter.editingFinished.connect(self.change_counts)
        counts = WrappingLayout()
        counts.addWidget(self.question_count)
        counts.addWidget(self.student_count)
        count_hint = QLabel("Lager aantal? Extra gegevens blijven bewaard.")
        count_hint.setWordWrap(True)
        counts.addWidget(count_hint)
        layout.addLayout(counts)
        self.undo_button = QPushButton("Ongedaan maken (Ctrl+Z)")
        self.undo_button.clicked.connect(self.undo)
        actions.addWidget(self.undo_button)
        self.direction = QComboBox()
        self.direction.addItem("Horizontaal (per leerling)", "horizontal")
        self.direction.addItem("Verticaal (per vraag)", "vertical")
        self.direction.addItem("Gecombineerd", "hybrid")
        actions.addWidget(self.direction)
        self.block_size = QSpinBox()
        self.block_size.setRange(1, 50)
        self.block_size.setValue(3)
        self.block_size.setPrefix("Gecombineerd: ")
        self.block_size.setSuffix(" vragen per blok")
        self.block_size.setVisible(False)
        self.direction.currentIndexChanged.connect(lambda: self.block_size.setVisible(self.direction.currentData() == "hybrid"))
        actions.addWidget(self.block_size)
        self.precision = QComboBox()
        self.precision.setToolTip("Hele punten: automatisch verder bij vragen met een maximum van hoogstens 9. "
                                 "Bij een hoger maximum bevestigt u met Enter of Tab. "
                                 "Bij decimalen bevestigt u met Enter of Tab.")
        for label, value in (("Hele punten", 1.0), ("Halve punten", .5), ("Kwartpunten", .25), ("Tiende punten", .1)):
            self.precision.addItem(label, value)
        self.precision.setCurrentIndex(max(0, self.precision.findData(float(database.meta("score_step", "1")))))
        self.precision.currentIndexChanged.connect(self.change_precision)
        actions.addWidget(self.precision)
        for label, fn in (("Vragen indelen", self.metadata), ("Meerkeuze instellen", self.multiple_choice),
                          ("Leerling verwijderen", self.remove_student)):
            button = QPushButton(label)
            button.clicked.connect(fn)
            actions.addWidget(button)
        layout.addLayout(actions)
        self.table = QTableWidget()
        from .results_page import ActiveStudentHeader, ActiveQuestionHeader
        self.table.setVerticalHeader(ActiveStudentHeader(self.table))
        self.table.setHorizontalHeader(ActiveQuestionHeader(self.table))
        self.table.setItemDelegate(ScoreEntryDelegate(self))
        self.table.currentCellChanged.connect(self.highlight_current)
        self.table.setAlternatingRowColors(False)
        self.table.cellChanged.connect(self.cell_changed)
        self.table.installEventFilter(self)
        self.table.setSelectionMode(QTableWidget.SelectionMode.ContiguousSelection)
        layout.addWidget(self.table, 1)
        self.legend = QLabel()
        self.legend.setWordWrap(True)
        layout.addWidget(self.legend)
        self.first_missing = QPushButton("Ga naar de eerste ontbrekende invoer")
        self.first_missing.clicked.connect(self.go_missing)
        layout.addWidget(self.first_missing)
        self.stack.addWidget(self.input)
        self.saved = QLabel()
        root.addWidget(self.saved)
        for sequence, callback in (("Ctrl+V", self.paste), ("Ctrl+C", self.copy), ("Ctrl+Z", self.undo), ("Ctrl+S", self.save)):
            shortcut = QShortcut(QKeySequence(sequence), self)
            shortcut.setContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
            shortcut.activated.connect(callback)
        self.refresh()

    @property
    def is_saved(self):
        return self.database.path.name != ":memory:"

    def remember(self, snapshot):
        self.history.append(snapshot)
        self.history = self.history[-25:]

    def snapshot(self):
        self.database.connection.commit()
        return self.database.connection.serialize()

    def restore(self, snapshot):
        self.database.connection.rollback()
        memory = sqlite3.connect(":memory:")
        try:
            memory.deserialize(snapshot)
            memory.backup(self.database.connection)
        finally:
            memory.close()

    def refresh(self):
        self.loading = True
        self.questions = active_questions(self.database)
        self.students = active_students(self.database)
        self.question_count.setValue(len(self.questions))
        self.student_count.setValue(len(self.students))
        attempts, scores, responses = stored_results(self.database, 1)
        self.table.clear()
        self.table.setColumnCount(len(self.questions) + 2)
        self.table.setRowCount(len(self.students) + 2)
        self.table.setHorizontalHeaderLabels(["Leerlingnaam"] + ["Vraag " + q["label"] for q in self.questions] + ["Totaal"])
        stored_students = int(self.database.scalar("SELECT COUNT(*) FROM test_students WHERE test_id=1"))
        hidden_students = stored_students > len(self.students)
        self.table.setVerticalHeaderLabels(["Maximum"] + [s["display_name"] for s in self.students] + ["Verborgen" if hidden_students else "+ Nieuw"])
        header_width = max(90, min(200, max((self.table.fontMetrics().horizontalAdvance(s["display_name"]) + 20 for s in self.students), default=90)))
        self.table.verticalHeader().setFixedWidth(header_width)
        for col, value in enumerate(["Maximaal aantal punten"] + [f"{q['maximum_score']:g}" if q["maximum_score"] else "" for q in self.questions] + [f"{sum(q['maximum_score'] for q in self.questions):g}"]):
            self.set_cell(0, col, value, "#dbeafe", editable=0 < col <= len(self.questions))
        for row, student in enumerate(self.students, 1):
            self.set_cell(row, 0, student["display_name"])
            for col, q in enumerate(self.questions, 1):
                value = scores.get((student["id"], q["id"]))
                response = responses.get((student["id"], q["id"]), "")
                raw = response if response or q["is_multiple_choice"] else "" if value is None else f"{value:g}"
                color, foreground = '#ffffff', '#1f3352'
                if raw == 'N':
                    color, foreground = '#fff8ed', '#9a6a1d'
                elif q['is_multiple_choice'] and raw and value is not None:
                    correct = value == q['maximum_score']
                    color, foreground = ('#e8f7e8', '#1f7a1f') if correct else ('#fdecec', '#a02828')
                self.set_cell(row, col, raw, color)
                self.table.item(row, col).setForeground(QColor(foreground))
            values = [scores.get((student["id"], q["id"])) for q in self.questions]
            total = sum(value for value in values if value is not None) if any(value is not None for value in values) else None
            complete = bool(values) and all(value is not None for value in values) and not any(r == row for r, _ in self.invalid)
            self.set_cell(row, len(self.questions) + 1, "" if total is None else f"{total:g}", "#dbeafe" if complete else "#edf2f7", False)
            total_cell = self.table.item(row, len(self.questions) + 1)
            total_cell.setForeground(QColor('#1d4ed8' if complete else '#1f3352'))
            font = total_cell.font()
            font.setBold(complete)
            total_cell.setFont(font)
        new = len(self.students) + 1
        for col in range(len(self.questions) + 2):
            self.set_cell(new, col, "Verhoog Aantal leerlingen voor de bewaarde rijen" if hidden_students and col == 0 else "", editable=not hidden_students and col <= len(self.questions))
        if hidden_students:
            self.table.setRowHeight(new, 60)
        self.table.setColumnWidth(0, 235)
        for col in range(1, len(self.questions) + 1):
            self.table.setColumnWidth(col, 95)
            q = self.questions[col - 1]
            self.table.horizontalHeaderItem(col).setToolTip("Antwoordletter invullen. Juiste antwoord: " + str(q["multiple_choice_answer"]) if q["is_multiple_choice"] else "Punten invullen, of N voor niet gemaakt.")
        for (row, col), (raw, message) in self.invalid.items():
            if row < self.table.rowCount() and col < self.table.columnCount():
                self.set_cell(row, col, raw)
                self.table.item(row, col).setToolTip(message)
        self.loading = False
        self.highlight_current(self.table.currentRow(), self.table.currentColumn())
        self.update_status()

    def highlight_current(self, row, col, previous_row=-1, previous_col=-1):
        if self.loading:
            return
        self.table.verticalHeader().set_active_row(row if 1 <= row <= len(self.students) else -1)
        self.table.horizontalHeader().set_active_column(col if 1 <= col <= len(self.questions) else -1)
        blocked = self.table.blockSignals(True)
        try:
            for r in {previous_row, row}:
                if 1 <= r <= len(self.students):
                    cell = self.table.item(r, 0)
                    if cell is not None:
                        cell.setBackground(QColor('#dbeafe' if r == row else '#ffffff'))
                        font = cell.font()
                        font.setBold(r == row)
                        cell.setFont(font)
        finally:
            self.table.blockSignals(blocked)
        self.table.viewport().update()

    def set_cell(self, row, col, value, color="#ffffff", editable=True):
        cell = QTableWidgetItem(str(value))
        cell.setBackground(QColor(color))
        if not editable:
            cell.setFlags(cell.flags() & ~Qt.ItemFlag.ItemIsEditable)
        if col:
            cell.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        self.table.setItem(row, col, cell)

    def update_status(self):
        from .norming import has_active_normalization
        state = completion(self.database)
        ready = state["ready"] and not self.invalid
        finalized = has_active_normalization(self.database, 1)
        self.norm_button.setEnabled(ready)
        self.analysis_button.setEnabled(ready)
        partial = ready and not finalized
        self.analysis_button.setProperty("partial_analysis", partial)
        self.analysis_button.setStyleSheet("QPushButton { background: #d97706; color: white; } QPushButton:hover { background: #b45309; }" if partial else "")
        if self.invalid:
            message = next(iter(self.invalid.values()))[1]
        elif state["maxima"]:
            message = f"Vul nog {len(state['maxima'])} maximumscores in."
        elif not state["students"]:
            message = "Voeg leerlingnamen toe of plak een blok uit Excel."
        elif state["missing"]:
            message = f"Nog {len(state['missing'])} scores invullen voordat Normeren en Analyseren beschikbaar zijn."
        elif not finalized:
            message = "Alle scores zijn ingevuld. U kunt de analyse al bekijken. Cijfers en cijferafhankelijke onderdelen komen beschikbaar na het vaststellen van de normering."
        else:
            message = "Alle scores zijn ingevuld. U kunt nu normeren en analyseren."
        self.status.setText(message)
        self.norm_button.setToolTip(message if not ready else "Kies de normering voor deze toets.")
        self.analysis_button.setToolTip(message if not ready or partial else "Bekijk de volledige analyse van deze toets, inclusief cijfers.")
        self.legend.setText("Kleurenlegenda: blauw = actieve leerling en vraag, maximumscores en volledig ingevuld totaal · "
                            "wit = gewone punten of leeg · groen = goed meerkeuzeantwoord · lichtrood = fout meerkeuzeantwoord · "
                            "lichtoranje = N, niet gemaakt · rode rand = ongeldige invoer. "
                            "Oranje analyseknop = normering nog niet vastgesteld.")
        self.first_missing.setEnabled(not ready)
        self.undo_button.setEnabled(bool(self.history) or bool(self.invalid))
        if self.invalid:
            self.saved.setText("Ongeldige invoer is nog niet opgeslagen. Corrigeer de cel met de rode rand of gebruik Ongedaan maken.")
        elif self.is_saved:
            self.saved.setText("Alles opgeslagen · " + str(self.database.path))
        else:
            self.saved.setText("Nog geen bestand gekozen. Klik op Opslaan om uw analyse te bewaren.")

    def apply_cell(self, row, col, raw):
        raw = str(raw).strip()
        if row == 0:
            if col == 0:
                if raw and raw.casefold() not in {"maximaal aantal punten", "maximum", "max. score:"}:
                    raise ValueError("Plak leerlingnamen vanaf de eerste leerlingregel onder de maximumscores.")
                return
            maximum = float(raw.replace(",", ".")) if raw else 0.0
            import math
            if not math.isfinite(maximum) or maximum <= 0:
                raise ValueError(f"Vul bij vraag {col} een maximumscore groter dan 0 in.")
            q = self.questions[col - 1]
            if not q["is_multiple_choice"] and self.database.scalar("SELECT COUNT(*) FROM scores WHERE question_id=? AND score>?", (q["id"], maximum)):
                raise ValueError("Er zijn al scores hoger dan dit maximum. Pas die eerst aan.")
            if q["maximum_score"] != maximum:
                self.database.connection.execute("UPDATE matrix_questions SET maximum_score=? WHERE id=?", (maximum, q["id"]))
                self.database.connection.execute("UPDATE normalizations SET is_active=0 WHERE test_id=1")
                q["maximum_score"] = maximum
                if q["is_multiple_choice"]:
                    for stored in self.database.rows("SELECT a.student_id,s.response_text FROM scores s JOIN test_attempts a ON a.id=s.attempt_id WHERE s.question_id=?", (q["id"],)):
                        save_score(self.database, 1, stored["student_id"], q["id"], stored["response_text"] or "", commit=False)
            return
        if row > 10000:
            raise ValueError("Er kunnen maximaal 10000 leerlingen worden ingevoerd.")
        if row > len(self.students):
            if self.database.scalar("SELECT COUNT(*) FROM test_students WHERE test_id=1") > len(self.students):
                raise ValueError("Er zijn bewaarde leerlingregels verborgen. Verhoog eerst Aantal leerlingen om die terug te zien.")
            set_counts(self.database, len(self.questions), row)
            self.students = active_students(self.database)
        student = self.students[row - 1]
        if col == 0:
            if not raw:
                raise ValueError("Vul een leerlingnaam in. Gebruik Leerling verwijderen om iemand weg te halen.")
            self.database.connection.execute("UPDATE students SET display_name=? WHERE id=?", (raw, student["id"]))
        else:
            q = self.questions[col - 1]
            if q["maximum_score"] <= 0:
                raise ValueError(f"Vul eerst de maximumscore van vraag {col} in.")
            try:
                save_score(self.database, 1, student["id"], q["id"], raw, commit=False,
                           score_step=float(self.database.meta("score_step", "1")))
            except ValueError as error:
                raise ValueError(f"{student['display_name']}, vraag {q['label']}: {error}") from error

    def cell_changed(self, row, col):
        if self._released or self.loading or col > len(self.questions) or (row == 0 and col == 0):
            return
        raw = self.table.item(row, col).text()
        if row > len(self.students) and not raw.strip():
            return
        self.edit_cells([(row, col, raw)])

    def edit_cells(self, entries, *, advance=True):
        before = self.snapshot()
        row = col = 0
        raw = ""
        try:
            with self.database.connection:
                for row, col, raw in entries:
                    self.apply_cell(row, col, raw)
                recalculate_grades(self.database)
        except (ValueError, sqlite3.Error) as error:
            self.restore(before)
            self.invalid[row, col] = (str(raw), str(error))
            self.refresh()
            self.table.setCurrentCell(row, col)
            return False
        self.remember(before)
        for row, col, _ in entries:
            self.invalid.pop((row, col), None)
        self.refresh()
        if entries and advance:
            self.advance(entries[-1][0], entries[-1][1])
        return True

    def is_score_cell(self, row, col):
        return 1 <= row <= len(self.students) and 1 <= col <= len(self.questions)

    def quick_score_complete(self, row, col, text):
        if not self.is_score_cell(row, col) or self.precision.currentData() != 1.0:
            return False
        q = self.questions[col - 1]
        if q['is_multiple_choice'] or not text.isascii() or not text.isdigit():
            return False
        value = int(text)
        maximum = float(q['maximum_score'] or 0)
        return 0 < maximum <= 9 and value <= maximum

    def clear_scores(self, cells=None):
        row, col = self.table.currentRow(), self.table.currentColumn()
        if cells is None:
            cells = [(i.row(), i.column()) for i in self.table.selectedIndexes()]
        entries = [(r, c, '') for r, c in cells if self.is_score_cell(r, c)]
        if entries:
            self.edit_cells(entries, advance=False)
            self.table.setCurrentCell(row, col)

    def advance(self, row, col):
        direction = self.direction.currentData()
        end = len(self.questions)
        if row == 0:
            target = (0, col + 1) if col < end else (1, 0)
        elif direction == "vertical" and col > 0:
            target = (row + 1, col) if row < len(self.students) else (1, min(end, col + 1))
        elif direction == "hybrid" and col > 0:
            block = self.block_size.value()
            block_start = ((col - 1) // block) * block + 1
            block_end = min(end, block_start + block - 1)
            target = (row, col + 1) if col < block_end else ((row + 1, block_start) if row < len(self.students) else (1, block_end + 1) if block_end < end else (row + 1, 0))
        else:
            target = (row, col + 1) if col < end else (row + 1, 0)
        self.table.setCurrentCell(min(target[0], self.table.rowCount() - 1), min(target[1], end))

    def eventFilter(self, watched, event):
        if watched is self.table and event.type() == QEvent.Type.KeyPress:
            if self.table.state() != QTableWidget.State.EditingState:
                if event.key() == Qt.Key.Key_Delete:
                    self.clear_scores()
                    return True
                row, col = self.table.currentRow(), self.table.currentColumn()
                if (not event.modifiers() & (Qt.KeyboardModifier.ControlModifier | Qt.KeyboardModifier.AltModifier | Qt.KeyboardModifier.MetaModifier)
                        and self.quick_score_complete(row, col, event.text())):
                    self.edit_cells([(row, col, event.text())])
                    return True
            if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter) and self.table.state() != QTableWidget.State.EditingState:
                self.advance(self.table.currentRow(), self.table.currentColumn())
                return True
        return super().eventFilter(watched, event)

    def paste(self):
        focus = QApplication.focusWidget()
        if isinstance(focus, QLineEdit) and not self.table.isAncestorOf(focus):
            focus.paste()
            return
        if self.stack.currentWidget() is not self.input:
            return
        rows = list(csv.reader(io.StringIO(QApplication.clipboard().text()), delimiter="\t"))
        if not rows:
            return
        row, col = max(0, self.table.currentRow()), max(0, self.table.currentColumn())
        if rows[0] and rows[0][0].strip() == "Leerlingnaam" and row == 0 and col == 0:
            rows = rows[1:]
        if len(rows) > 10002 or any(len(values) > len(self.questions) + 2 - col for values in rows):
            QMessageBox.warning(self, "Plakken", "Dit blok past niet in de tabel. Kies de juiste begincel en kopieer alleen de gewenste kolommen.")
            return
        entries = []
        for r, values in enumerate(rows, row):
            for c, value in enumerate(values, col):
                if c <= len(self.questions) and not (r > len(self.students) and not value.strip()):
                    entries.append((r, c, value))
        self.edit_cells(entries)

    def copy(self):
        focus = QApplication.focusWidget()
        if isinstance(focus, QLineEdit) and not self.table.isAncestorOf(focus):
            focus.copy()
            return
        selected = self.table.selectedRanges()
        if not selected:
            return
        area = selected[0]
        values = [[self.table.item(r, c).text() if self.table.item(r, c) else "" for c in range(area.leftColumn(), area.rightColumn() + 1)] for r in range(area.topRow(), area.bottomRow() + 1)]
        output = io.StringIO()
        csv.writer(output, delimiter="\t", lineterminator="\n").writerows(values)
        QApplication.clipboard().setText(output.getvalue())

    def undo(self):
        if self.stack.currentWidget() is not self.input:
            return
        if self.invalid:
            self.invalid.clear()
            self.refresh()
            return
        if self.history:
            self.restore(self.history.pop())
        self.invalid.clear()
        self.name.setText(str(self.database.scalar("SELECT name FROM tests WHERE id=1")))
        self.precision.blockSignals(True)
        self.precision.setCurrentIndex(max(0, self.precision.findData(float(self.database.meta("score_step", "1")))))
        self.precision.blockSignals(False)
        self.refresh()

    def go_missing(self):
        if self.invalid:
            row, col = next(iter(self.invalid))
        else:
            state = completion(self.database)
            if state["maxima"]:
                row, col = 0, next(i for i, q in enumerate(self.questions, 1) if q["id"] == state["maxima"][0])
            elif state["missing"]:
                sid, qid = state["missing"][0]
                row = next(i for i, s in enumerate(self.students, 1) if s["id"] == sid)
                col = next(i for i, q in enumerate(self.questions, 1) if q["id"] == qid)
            else:
                row, col = len(self.students) + 1, 0
        self.table.setCurrentCell(row, col)
        self.table.scrollToItem(self.table.item(row, col))
        self.table.setFocus()

    def rename(self):
        if self._released:
            return
        name = self.name.text().strip()
        if not name:
            self.name.setText(str(self.database.scalar("SELECT name FROM tests WHERE id=1")))
            return
        if name and name != self.database.scalar("SELECT name FROM tests WHERE id=1"):
            self.remember(self.snapshot())
            with self.database.connection:
                self.database.connection.execute("UPDATE tests SET name=? WHERE id=1", (name,))
                self.database.set_meta("subject_name", name)
            self.update_status()

    def change_precision(self):
        if self._released:
            return
        before = self.snapshot()
        step = self.precision.currentData()
        values = [float(row[0]) for row in self.database.rows("SELECT score FROM scores WHERE score IS NOT NULL")]
        import math
        if any(not math.isclose(value / step, round(value / step), abs_tol=1e-8) for value in values):
            QMessageBox.warning(self, "Puntendetail", "Er zijn scores die niet bij deze puntstap passen. Pas die scores eerst aan.")
            self.precision.blockSignals(True)
            self.precision.setCurrentIndex(self.precision.findData(float(self.database.meta("score_step", "1"))))
            self.precision.blockSignals(False)
            return
        self.database.set_meta("score_step", str(step))
        self.database.connection.commit()
        self.remember(before)
        self.update_status()

    def change_counts(self):
        if self._released or self.loading:
            return
        questions, students = self.question_count.value(), self.student_count.value()
        if (questions, students) == (len(self.questions), len(self.students)):
            return
        before = self.snapshot()
        with self.database.connection:
            set_counts(self.database, questions, students)
            recalculate_grades(self.database)
        self.remember(before)
        self.invalid.clear()
        self.refresh()

    def metadata(self):
        dialog = MetadataDialog(self.database, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            before = self.snapshot()
            try:
                with self.database.connection:
                    dialog.apply()
            except Exception as error:
                QMessageBox.warning(self, "Niet opgeslagen", str(error))
                return
            self.remember(before)
            self.update_status()

    def multiple_choice(self):
        dialog = MultipleChoiceDialog(self.database, self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        before = self.snapshot()
        try:
            with self.database.connection:
                for q in dialog.questions:
                    check, combo = dialog.controls[q["id"]]
                    is_mc, answer = check.isChecked(), combo.currentText() or None
                    if is_mc != bool(q["is_multiple_choice"]) and self.database.scalar("SELECT COUNT(*) FROM scores WHERE question_id=?", (q["id"],)):
                        raise ValueError("Gebruik eerst Ongedaan maken of wis de scores van deze vraag voordat u wisselt tussen punten en antwoordletters.")
                    self.database.connection.execute("UPDATE matrix_questions SET is_multiple_choice=?,multiple_choice_answer=? WHERE id=?", (int(is_mc), answer if is_mc else None, q["id"]))
                    if is_mc:
                        for row in self.database.rows("SELECT a.student_id,s.response_text FROM scores s JOIN test_attempts a ON a.id=s.attempt_id WHERE s.question_id=?", (q["id"],)):
                            save_score(self.database, 1, row["student_id"], q["id"], row["response_text"] or "", commit=False)
                recalculate_grades(self.database)
        except ValueError as error:
            QMessageBox.warning(self, "Meerkeuze instellen", str(error))
            return
        self.remember(before)
        self.refresh()

    def remove_student(self):
        row = self.table.currentRow()
        if not 1 <= row <= len(self.students):
            return
        student = self.students[row - 1]
        if question(self, "Leerling verwijderen", f"{student['display_name']} en de bijbehorende scores verwijderen?") != QMessageBox.StandardButton.Yes:
            return
        before = self.snapshot()
        with self.database.connection:
            self.database.connection.execute("DELETE FROM students WHERE id=?", (student["id"],))
            self.database.set_meta("simple_student_count", str(len(self.students) - 1))
        self.remember(before)
        self.invalid.clear()
        self.refresh()

    def save(self, choose=False):
        self.commit_input_editor()
        self.rename()
        if self.invalid:
            QMessageBox.warning(self, "Invoer corrigeren", "Corrigeer eerst de invoercellen met een rode rand. De overige wijzigingen zijn wel bewaard als u al een bestand gekozen heeft.")
            return False
        if choose or not self.is_saved:
            path, _ = QFileDialog.getSaveFileName(self, "Analyse opslaan", self.name.text().strip() + SIMPLE_SUFFIX, "Toetsvizier-analyse (*.tviz)")
            if not path:
                return False
            if not path.lower().endswith(SIMPLE_SUFFIX):
                path += SIMPLE_SUFFIX
            try:
                save_as(self.database, path)
            except Exception as error:
                QMessageBox.warning(self, "Niet opgeslagen", str(error))
                return False
        self.database.connection.commit()
        self.update_status()
        return True

    def confirm_close(self):
        self.commit_input_editor()
        self.rename()
        if self.invalid:
            if question(self, "Ongeldige invoer", "De invoer met een rode rand is niet opgeslagen. Toch sluiten?", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No) != QMessageBox.StandardButton.Yes:
                return False
        if not self.is_saved:
            answer = question(self, "Analyse bewaren", "Wilt u deze analyse opslaan voordat u sluit?", QMessageBox.StandardButton.Save | QMessageBox.StandardButton.Discard | QMessageBox.StandardButton.Cancel, QMessageBox.StandardButton.Save)
            if answer == QMessageBox.StandardButton.Cancel:
                return False
            if answer == QMessageBox.StandardButton.Save:
                return self.save()
        return True

    def commit_input_editor(self):
        editor = QApplication.focusWidget()
        if isinstance(editor, QLineEdit) and self.table.isAncestorOf(editor):
            delegate = self.table.itemDelegate()
            delegate.commitData.emit(editor)
            try:
                delegate.closeEditor.emit(editor, QAbstractItemDelegate.EndEditHint.NoHint)
            except RuntimeError:
                pass

    def request_close(self):
        if self.confirm_close():
            self.back_to_start.emit()

    def show_intro(self, force=False):
        if not force and str(self.settings.value("simple/skip_intro", "0")) == "1":
            return
        dialog = InputIntroduction(self)
        dialog.skip.setChecked(str(self.settings.value("simple/skip_intro", "0")) == "1")
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.settings.setValue("simple/skip_intro", "1" if dialog.skip.isChecked() else "0")

    def show_input(self):
        self.stack.setCurrentWidget(self.input)
        self.refresh()

    def release_views(self):
        self._released = True
        if self.active_page is not None:
            self.stack.removeWidget(self.active_page)
            self.active_page.deleteLater()
            self.active_page = None
        if self.projection is not None:
            self.projection.close()
            self.projection = None

    def open_module(self, kind):
        if not completion(self.database)["ready"] or self.invalid:
            self.go_missing()
            return
        # Recreate on opening: all calculations use the latest table and metadata.
        if self.active_page is not None:
            self.stack.removeWidget(self.active_page)
            self.active_page.deleteLater()
        if self.projection is not None:
            self.projection.close()
        self.projection = project_active(self.database)
        from .app import NormingPage
        from .analysis_page import AnalysisPage
        page = NormingPage(self.projection, 1) if kind == "norm" else AnalysisPage(self.projection, 1)
        if kind == "norm":
            page.normalization_changed.connect(self.update_status)
        for control in (page.level_filter, page.grade_filter, page.test):
            control.hide()
        for label in page.findChildren(QLabel):
            if label.text() in {"Niveau", "Jaarlaag", "Toets"}:
                label.hide()
        for frame in page.findChildren(QFrame):
            if frame.objectName() == "filterCard":
                frame.hide()
        page.select_test(1)
        self.active_page = page
        self.stack.addWidget(page)
        self.stack.setCurrentWidget(page)

    def export_excel(self):
        self.commit_input_editor()
        if self.invalid:
            QMessageBox.warning(self, "Invoer corrigeren", "Corrigeer eerst de cellen met een rode rand voordat u exporteert.")
            return
        path, _ = QFileDialog.getSaveFileName(self, "Exporteren naar Excel", self.name.text() + ".xlsx", "Excelbestand (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            export_simple_excel(self.database, path)
        except Exception as error:
            QMessageBox.warning(self, "Export mislukt", str(error))
            return
        self.status.setText("Excelbestand opgeslagen: " + path)


def export_simple_excel(database, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scores"
    questions = active_questions(database)
    attempts, scores, responses = stored_results(database, 1)
    sheet.append(["Leerlingnaam"] + [q["label"] for q in questions] + ["Totaal"])
    sheet.append(["Maximaal aantal punten"] + [q["maximum_score"] for q in questions] + [sum(q["maximum_score"] for q in questions)])
    for student in active_students(database):
        values = [responses.get((student["id"], q["id"])) or scores.get((student["id"], q["id"])) for q in questions]
        numeric = [scores.get((student["id"], q["id"])) for q in questions]
        total = sum(value for value in numeric if value is not None) if any(value is not None for value in numeric) else None
        sheet.append([student["display_name"]] + values + [total])
    sheet.freeze_panes = "B3"
    sheet.column_dimensions["A"].width = 30
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="2463EB")
        cell.font = Font(color="FFFFFF", bold=True)
    keys = workbook.create_sheet("Vragen en metadata")
    definitions = list(database.rows("SELECT * FROM taxonomy_definitions ORDER BY id"))
    keys.append(["Vraag", "Maximumscore", "Antwoordsleutel"] + [d["name"] for d in definitions])
    for q in questions:
        categories = [database.scalar("SELECT v.name FROM question_taxonomy_values qtv JOIN taxonomy_values v ON v.id=qtv.taxonomy_value_id WHERE qtv.question_id=? AND qtv.taxonomy_id=?", (q["id"], d["id"])) for d in definitions]
        keys.append([q["label"], q["maximum_score"], q["multiple_choice_answer"] if q["is_multiple_choice"] else None] + categories)
    workbook.save(path)
    workbook.close()
