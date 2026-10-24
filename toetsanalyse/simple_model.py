"""Single-test documents, separate from subject databases in the extended mode."""
from __future__ import annotations

import math
import sqlite3
from contextlib import closing
from pathlib import Path

from .database import SubjectDatabase
from .results import save_score, stored_results, test_questions, test_students

SIMPLE_SUFFIX = ".tviz"


def create_analysis(question_count: int, name: str = "Nieuwe analyse", student_count: int = 0) -> SubjectDatabase:
    if not 1 <= question_count <= 1000:
        raise ValueError("Kies tussen 1 en 1000 vragen.")
    database = SubjectDatabase.create(":memory:", name, "Losse analyse")
    database.set_meta("document_kind", "simple_analysis")
    database.set_meta("score_step", "1")
    database.set_meta("simple_question_count", str(question_count))
    database.set_meta("simple_student_count", "0")
    database.connection.execute("INSERT INTO tests(school_year_id,name,period,test_type) VALUES(1,?,'','Losse toets')", (name,))
    database.connection.executemany("INSERT INTO matrix_questions(test_id,question_number,maximum_score) VALUES(1,?,0)",
                                    [(str(n),) for n in range(1, question_count + 1)])
    set_counts(database, question_count, student_count)
    database.connection.commit()
    return database


def open_analysis(path: str | Path) -> SubjectDatabase:
    path = Path(path)
    # Verify before migrating, so an accidentally selected subject file is unchanged.
    with closing(sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)) as connection:
        row = connection.execute("SELECT value FROM app_meta WHERE key='document_kind'").fetchone()
        if not row or row[0] != "simple_analysis":
            raise ValueError("Dit is geen opgeslagen simpele analyse. Open vakbestanden via Uitgebreid.")
    return SubjectDatabase.open(path)


def save_as(database: SubjectDatabase, path: str | Path) -> None:
    path = Path(path)
    if path.resolve() == database.path.resolve():
        database.connection.commit()
        return
    database.connection.commit()
    connection = sqlite3.connect(path)
    try:
        database.connection.backup(connection)
    except Exception:
        connection.close()
        raise
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys=ON")
    database.connection.close()
    database.connection = connection
    database.path = path


def add_student(database: SubjectDatabase, name: str) -> int:
    name = name.strip()
    if not name:
        raise ValueError("Vul een leerlingnaam in.")
    sid = database.connection.execute("INSERT INTO students(display_name) VALUES(?)", (name,)).lastrowid
    database.connection.execute("INSERT INTO test_students(test_id,student_id) VALUES(1,?)", (sid,))
    database.set_meta("simple_student_count", str(database.scalar("SELECT COUNT(*) FROM test_students WHERE test_id=1")))
    return int(sid)


def active_questions(database):
    count = int(database.meta("simple_question_count", "1000"))
    ids = [row[0] for row in database.rows("SELECT id FROM matrix_questions WHERE test_id=1 ORDER BY id LIMIT ?", (count,))]
    lookup = {q["id"]: q for q in test_questions(database, 1)}
    return [lookup[qid] for qid in ids]


def active_students(database):
    count = int(database.meta("simple_student_count", "10000"))
    return [dict(row) for row in database.rows("SELECT s.id,s.display_name FROM students s JOIN test_students ts ON ts.student_id=s.id WHERE ts.test_id=1 ORDER BY s.id LIMIT ?", (count,))]


def active_taxonomies(database):
    selected = database.rows('SELECT d.* FROM taxonomy_definitions d JOIN test_taxonomy_selections s '
                             'ON s.taxonomy_id=d.id WHERE s.test_id=1 ORDER BY d.id')
    if selected or database.meta('simple_taxonomy_selection_configured', '0') == '1':
        return selected
    # Older documents can contain classifications without a saved selection.
    return database.rows('SELECT DISTINCT d.* FROM taxonomy_definitions d '
                         'JOIN question_taxonomy_values q ON q.taxonomy_id=d.id '
                         'JOIN matrix_questions mq ON mq.id=q.question_id WHERE mq.test_id=1 ORDER BY d.id')


def set_counts(database, question_count, student_count):
    if not 1 <= question_count <= 1000 or not 0 <= student_count <= 10000:
        raise ValueError("Kies 1–1000 vragen en 0–10000 leerlingen.")
    existing_q = int(database.scalar("SELECT COUNT(*) FROM matrix_questions WHERE test_id=1"))
    existing_s = int(database.scalar("SELECT COUNT(*) FROM test_students WHERE test_id=1"))
    for number in range(existing_q + 1, question_count + 1):
        database.connection.execute("INSERT INTO matrix_questions(test_id,question_number,maximum_score) VALUES(1,?,0)", (str(number),))
    for number in range(existing_s + 1, student_count + 1):
        add_student(database, f"Leerling {number}")
    if question_count != int(database.meta("simple_question_count", str(existing_q))):
        database.connection.execute("UPDATE normalizations SET is_active=0 WHERE test_id=1")
        invalidate_grades(database)
    database.set_meta("simple_question_count", str(question_count))
    database.set_meta("simple_student_count", str(student_count))


def completion(database: SubjectDatabase) -> dict:
    questions = active_questions(database)
    students = active_students(database)
    _, scores, _ = stored_results(database, 1)
    missing = [(student["id"], q["id"]) for student in students for q in questions if (student["id"], q["id"]) not in scores]
    maxima = [q["id"] for q in questions if float(q["maximum_score"]) <= 0]
    keys = [q["id"] for q in questions if q["is_multiple_choice"] and not q["multiple_choice_answer"]]
    return {"ready": bool(students) and bool(questions) and not missing and not maxima and not keys,
            "missing": missing, "maxima": maxima, "keys": keys, "students": len(students), "questions": len(questions)}


def invalidate_grades(database: SubjectDatabase) -> None:
    database.connection.execute("UPDATE test_attempts SET grade=NULL WHERE test_id=1")


def recalculate_grades(database: SubjectDatabase) -> None:
    from .norming import calculate_grade, has_active_normalization, load_normalization, maximum_score
    invalidate_grades(database)
    if not completion(database)["ready"] or not has_active_normalization(database, 1):
        return
    settings = load_normalization(database, 1)
    maximum = sum(q["maximum_score"] for q in active_questions(database))
    _, scores, _ = stored_results(database, 1)
    for student in active_students(database):
        total = sum(scores.get((student["id"], q["id"]), 0) for q in active_questions(database))
        database.connection.execute("UPDATE test_attempts SET grade=? WHERE test_id=1 AND student_id=?", (calculate_grade(total, maximum, settings), student["id"]))


def project_active(database):
    database.connection.commit()
    projected = SubjectDatabase(":memory:")
    database.connection.backup(projected.connection)
    connection = projected.connection
    connection.execute("PRAGMA foreign_keys=OFF")
    taxonomies = {row['id'] for row in active_taxonomies(database)}
    for row in projected.rows('SELECT DISTINCT taxonomy_id FROM question_taxonomy_values'):
        if row['taxonomy_id'] not in taxonomies:
            connection.execute('DELETE FROM question_taxonomy_values WHERE taxonomy_id=?', (row['taxonomy_id'],))
    questions = {q["id"] for q in active_questions(database)}
    students = {s["id"] for s in active_students(database)}
    for row in projected.rows("SELECT id FROM matrix_questions WHERE test_id=1"):
        if row["id"] not in questions:
            connection.execute("DELETE FROM scores WHERE question_id=?", (row["id"],))
            connection.execute("DELETE FROM matrix_questions WHERE id=?", (row["id"],))
    for row in projected.rows("SELECT id,student_id FROM test_attempts WHERE test_id=1"):
        if row["student_id"] not in students:
            connection.execute("DELETE FROM test_attempts WHERE id=?", (row["id"],))
        else:
            connection.execute("UPDATE test_attempts SET total_score=(SELECT SUM(score) FROM scores WHERE attempt_id=?) WHERE id=?", (row["id"], row["id"]))
    connection.execute("DELETE FROM test_students WHERE test_id=1")
    connection.executemany("INSERT INTO test_students(test_id,student_id) VALUES(1,?)", [(sid,) for sid in students])
    connection.commit()
    projected.simple_source = database
    return projected


def persist_norming(projected):
    source = getattr(projected, "simple_source", None)
    if source is None:
        return
    with source.connection as connection:
        connection.execute("UPDATE normalizations SET is_active=0 WHERE test_id=1")
        for row in projected.rows("SELECT * FROM normalizations WHERE test_id=1 AND is_active=1"):
            connection.execute("INSERT INTO normalizations(test_id,method,configuration_json,decimal_places,rounding_method,is_active) VALUES(1,?,?,?,?,1)",
                               (row["method"], row["configuration_json"], row["decimal_places"], row["rounding_method"]))
        for row in projected.rows("SELECT id,grade FROM test_attempts WHERE test_id=1"):
            connection.execute("UPDATE test_attempts SET grade=? WHERE id=?", (row["grade"], row["id"]))


def import_cito(path: str | Path) -> SubjectDatabase:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    database = None
    try:
        if not {"Referentie", "Gegevens"}.issubset(workbook.sheetnames):
            raise ValueError("Dit bestand heeft niet de herkenbare CITO-tabbladen Referentie en Gegevens.")
        count = workbook["Referentie"]["C8"].value
        if not isinstance(count, (int, float)) or not math.isfinite(count) or int(count) != count or not 1 <= count <= 1000:
            raise ValueError("Het aantal vragen in Referentie ontbreekt of is ongeldig.")
        count = int(count)
        sheet = workbook["Gegevens"]
        rows = list(sheet.iter_rows(max_col=5 + count, values_only=True))
        if len(rows) < 6 or str(rows[3][4]).strip().casefold() != "max. score:":
            raise ValueError("De indeling van dit CITO-bestand wordt niet herkend: de rij met maximumscores ontbreekt.")
        database = create_analysis(count, Path(path).stem)
        step_name = str(workbook["Referentie"]["C10"].value or "hele").strip().casefold()
        step = {"hele": 1, "halve": .5, "kwart": .25, "tiende": .1}.get(step_name)
        if step is None:
            raise ValueError("Het puntendetailniveau in Referentie wordt niet herkend.")
        database.set_meta("score_step", str(step))
        questions = test_questions(database, 1)
        for offset, question in enumerate(questions):
            maximum = rows[3][5 + offset]
            if not isinstance(maximum, (int, float)) or isinstance(maximum, bool) or not math.isfinite(maximum) or maximum <= 0:
                raise ValueError(f"De maximumscore van vraag {offset + 1} ontbreekt of is ongeldig.")
            label = rows[2][5 + offset]
            if label is None:
                label = rows[1][5 + offset] or str(offset + 1)
            database.connection.execute("UPDATE matrix_questions SET maximum_score=?,question_number=? WHERE id=?", (maximum, str(label), question["id"]))
        student_count = 0
        for row_number, row in enumerate(rows[5:], 6):
            if str(row[1] or "").strip().casefold().startswith(("gem. toetsscore", "cronbach")):
                break
            values = row[5:5 + count]
            if all(value is None or value == "" for value in values):
                continue  # Empty template rows and absent students are not imported.
            student_count += 1
            sid = add_student(database, str(row[1] or f"Leerling {student_count}"))
            for question, value in zip(questions, values):
                if value is None or value == "":
                    continue
                try:
                    save_score(database, 1, sid, question["id"], str(value), commit=False, score_step=step)
                except ValueError as error:
                    raise ValueError(f"Rij {row_number}, vraag {question['label']}: {error}") from error
        if not student_count:
            raise ValueError("Dit bestand bevat nog geen ingevulde leerlingenscores.")
        database.set_meta("source_cito", str(Path(path)))
        database.connection.commit()
        return database
    except Exception:
        if database is not None:
            database.close()
        raise
    finally:
        workbook.close()
